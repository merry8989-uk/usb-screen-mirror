#!/usr/bin/env python3
"""
PhoneMirror  -  Screen Mirror for Android, powered by scrcpy.

Ek Windows + Python GUI jo scrcpy ko wrap karta hai.
Features:
  * USB ya Wi-Fi (TCP/IP) connection
  * Multi-device: ek saath do ya zyada phone (har ek ka apna session)
  * Screen recording: session ko mp4 file me save
  * Camera mirroring: phone ka camera PC par (scrcpy 2.2+)
  * Low-latency presets, clear audio (Opus), input control
  * Auto-reconnect per device
  * Diagnostics: structured event logging (ProGallery Logging Spec v1.0) +
    logs export (.txt/.md/.json/.jsonl/.csv/.xlsx), email compose (pre-filled),
    copy, refresh, clean. On/off setting (default ON).

Yeh scrcpy ka engine use karta hai (hardware encode + USB transport).

Sirf Python 3 + scrcpy + adb chahiye. Tkinter Python ke saath aata hai.
(.xlsx export ke liye optional: pip install openpyxl)

Run:  python PhoneMirror.py
"""

import os
import re
import sys
import csv
import json
import time
import random
import hashlib
import shutil
import platform
import datetime
import subprocess
import threading
import queue
import webbrowser
import traceback
import tkinter as tk
from tkinter import ttk, messagebox, filedialog
from urllib.parse import quote

APP_TITLE = "PhoneMirror  -  Screen Mirror (scrcpy wrapper)"
APP_VERSION = "1.10.0"
DIAG_EMAIL = "br19shc39@gmail.com"

IS_WINDOWS = (os.name == "nt")
CREATE_NO_WINDOW = 0x08000000 if IS_WINDOWS else 0

CONFIG_DIR = os.path.join(os.path.expanduser("~"), ".phonemirror")
CONFIG_PATH = os.path.join(CONFIG_DIR, "config.json")


# --------------------------------------------------------------------------- #
#  Config persistence (diagnostics on/off yaad rakhne ke liye)
# --------------------------------------------------------------------------- #
def load_config():
    try:
        with open(CONFIG_PATH, encoding="utf-8") as f:
            data = json.load(f)
        return data if isinstance(data, dict) else {}
    except Exception:  # noqa: BLE001
        return {}


def save_config(cfg):
    try:
        os.makedirs(CONFIG_DIR, exist_ok=True)
        with open(CONFIG_PATH, "w", encoding="utf-8") as f:
            json.dump(cfg, f, indent=2)
    except Exception:  # noqa: BLE001
        pass


# --------------------------------------------------------------------------- #
#  Structured logging  -  ProGallery spec + log taxonomy
#  (common envelope, 12 categories, 6 severities, sampling, retention, PII redaction)
# --------------------------------------------------------------------------- #
LOG_SCHEMA_VERSION = "1.0.0"
LOG_MAX_EVENTS = 5000           # retention: memory/export me max events
LOG_RETENTION_DAYS = 7          # retention: isse purane events drop
LOG_SAMPLE_INFO_RATE = 1.0      # trace/debug/info sample rate (warn/error/fatal always 100%)
LOG_RELEASE_CHANNEL = "production"
APP_ID = "phonemirror-desktop"
LOG_ENVIRONMENT = "production"

# Log taxonomy: 12 categories
LOG_CATEGORIES = (
    "application", "error", "security", "performance", "user_activity",
    "infrastructure", "api_network", "database", "audit", "analytics",
    "background_job", "debug_trace",
)
# 6 severity levels
LOG_SEVERITIES = ("trace", "debug", "info", "warn", "error", "fatal")

# event_name -> category (log taxonomy)
EVENT_CATEGORY = {
    "session_start": "application",
    "session_end": "application",
    "mirror_result": "application",
    "device_state_resolved": "application",
    "device_refresh_result": "application",
    "screen_view": "user_activity",
    "mirror_requested": "user_activity",
    "preset_applied": "user_activity",
    "device_authorization_failed": "security",
    "adb_command": "api_network",
    "performance_marker": "performance",
    "export_saved": "audit",
    "email_compose": "audit",
    "share_bundle": "audit",
    "settings_changed": "audit",
    "diagnostics_toggled": "audit",
    "diagnostics_cleaned": "audit",
    "reconnect_attempt": "background_job",
    "reconnect_result": "background_job",
    "reconnect_loop": "background_job",
    "tool_probe": "background_job",
    "environment_snapshot": "infrastructure",
    "error_state_shown": "error",
    "crash_report": "error",
    "queue_drain": "debug_trace",
    "analytics_snapshot": "analytics",
}


def category_for(event_name, override=""):
    """Event ka log category nikalo (override pehle, warna map, warna application)."""
    if override:
        return override
    return EVENT_CATEGORY.get(event_name, "application")

_EMAIL_RE = re.compile(r"[A-Za-z0-9._%+\-]+@[A-Za-z0-9.\-]+\.[A-Za-z]{2,}")
_WINPATH_RE = re.compile(r"[A-Za-z]:\\[^\s\"]+")
_POSIXPATH_RE = re.compile(r"(?<![\w:])/(?:Users|home|tmp|var|mnt|media|sdcard|opt)/[^\s\"]*")


def redact(text):
    """PII redaction - app aur web ke liye identical rules.
    Emails aur file paths hataata hai (spec Part F)."""
    if text is None:
        return ""
    s = str(text)
    s = _EMAIL_RE.sub("<email>", s)
    s = _WINPATH_RE.sub("<path>", s)
    s = _POSIXPATH_RE.sub("<path>", s)
    return s


def salted_hash(value, salt):
    """user_id / device id sirf salted hash ke roop me (raw identifier nahi)."""
    if not value:
        return ""
    return "id_" + hashlib.sha256((salt + str(value)).encode("utf-8")).hexdigest()[:12]


def utc_now():
    return datetime.datetime.now(datetime.timezone.utc).strftime("%Y-%m-%dT%H:%M:%SZ")


def _rand_hex(n=16):
    return "".join(random.choice("0123456789abcdef") for _ in range(n))


def _platform_enum():
    s = platform.system().lower()
    if s.startswith("win"):
        return "windows"
    if s == "darwin":
        return "macos"
    return s or "unknown"


def _env_os():
    return "%s %s" % (platform.system(), platform.release())


def _device_model():
    return platform.machine() or platform.node() or "unknown"


def make_event(event_name, log_level, screen_or_route, session_id,
               attempt_id=None, user_id=None, network_type=None, props=None,
               category="application", message="", request_id=""):
    """Common event envelope: ProGallery spec + log taxonomy (category, severity)."""
    ev = {
        "event_name": event_name,
        "schema_version": LOG_SCHEMA_VERSION,
        "timestamp": utc_now(),
        "session_id": session_id,
        "platform": _platform_enum(),
        "app_id": APP_ID,
        "app_version": APP_VERSION,
        "environment": LOG_ENVIRONMENT,
        "os": _env_os(),
        "device_model": _device_model(),
        "screen_or_route": screen_or_route,
        "log_level": log_level,
        "severity": log_level,
        "category": category if category in LOG_CATEGORIES else "application",
        "message": redact(message or event_name),
        "release_channel": LOG_RELEASE_CHANNEL,
    }
    if attempt_id:
        ev["attempt_id"] = attempt_id
    if request_id:
        ev["request_id"] = request_id
    if user_id:
        ev["user_id"] = user_id
    if network_type:
        ev["network_type"] = network_type
    if props:
        for k, v in props.items():
            ev[k] = redact(v) if isinstance(v, str) else v
    return ev


class StructuredLogger:
    """Event envelope + log levels + sampling + retention."""

    def __init__(self, session_id, user_id=""):
        self.session_id = session_id
        self.user_id = user_id
        self.events = []
        self.attempt_seq = 0

    def new_attempt_id(self, prefix="att"):
        self.attempt_seq += 1
        return "%s_%04d" % (prefix, self.attempt_seq)

    def log(self, event_name, log_level, screen_or_route,
            attempt_id=None, props=None, network_type=None,
            category="application", message="", request_id=""):
        # sampling: warn/error/fatal hamesha 100%, trace/debug/info sampled
        if log_level in ("trace", "debug", "info") and LOG_SAMPLE_INFO_RATE < 1.0:
            if random.random() > LOG_SAMPLE_INFO_RATE:
                return None
        ev = make_event(event_name, log_level, screen_or_route, self.session_id,
                        attempt_id=attempt_id, user_id=self.user_id,
                        network_type=network_type, props=props,
                        category=category, message=message, request_id=request_id)
        self.events.append(ev)
        self._enforce_retention()
        return ev

    def _enforce_retention(self):
        if len(self.events) > LOG_MAX_EVENTS:
            del self.events[:len(self.events) - LOG_MAX_EVENTS]
        cutoff = datetime.datetime.now(datetime.timezone.utc) - datetime.timedelta(days=LOG_RETENTION_DAYS)
        keep = []
        for e in self.events:
            try:
                ts = datetime.datetime.strptime(e["timestamp"], "%Y-%m-%dT%H:%M:%SZ")
                ts = ts.replace(tzinfo=datetime.timezone.utc)
                if ts >= cutoff:
                    keep.append(e)
            except Exception:  # noqa: BLE001
                keep.append(e)
        self.events = keep

    def to_jsonl(self):
        return "\n".join(json.dumps(e, ensure_ascii=False) for e in self.events)


def sample_events(session_id="sess_sample"):
    """CI validation ke liye sample events (schema ke against check hote hain)."""
    lg = StructuredLogger(session_id)
    def e(name, level, screen, **kw):
        lg.log(name, level, screen, category=category_for(name), **kw)

    # application
    e("session_start", "info", "Main", props={"start_type": "cold"})
    e("session_end", "info", "Main", props={"reason": "user_close"})
    # infrastructure
    e("environment_snapshot", "info", "Main", props={"os": "Windows 11", "python": "3.12"})
    # user_activity
    e("screen_view", "info", "Diagnostics", props={"screen_name": "Diagnostics"})
    e("preset_applied", "info", "Main", props={"setting": "preset", "value": "Balanced"})
    # application (device)
    e("device_state_resolved", "info", "Main", props={"state": "authorized", "device_count": 1})
    # security
    e("device_authorization_failed", "warn", "Main", props={"state": "unauthorized", "device_count": 1})
    # user_activity (mirror)
    e("mirror_requested", "info", "Main", attempt_id="att_0001", props={"device_count": 1})
    e("mirror_result", "info", "Main", attempt_id="att_0001", props={"outcome": "success"})
    # api_network
    e("adb_command", "debug", "Main", props={"command": "adb devices", "outcome": "success", "latency_ms": 12})
    # performance
    e("performance_marker", "info", "Main", props={"flow": "adb_command", "duration_ms": 12})
    # background_job
    e("reconnect_attempt", "info", "Main", props={"attempt_number": 1})
    e("tool_probe", "debug", "Main", props={"outcome": "success"})
    # audit
    e("settings_changed", "info", "Main", props={"setting": "preset"})
    e("export_saved", "info", "Diagnostics", props={"format": "jsonl", "outcome": "success"})
    e("share_bundle", "info", "Diagnostics", props={"outcome": "success", "format": "zip"})
    # error
    e("error_state_shown", "error", "Main", props={"error_code": "scrcpy_missing"})
    e("crash_report", "fatal", "Main", props={"fatal": True, "exception_type": "RuntimeError",
                                             "stack_trace_hash": "abc123", "report_file": "crash-x.txt"})
    # analytics
    e("analytics_snapshot", "info", "Main", props={"mirrors_started": 3, "exports": 1})
    # debug_trace
    e("queue_drain", "trace", "Main", props={"processed": 250})
    return lg.events


class FileLogWriter:
    """Logs device par continuously save karta hai (JSONL events + plain log).
    Daily rotate + retention (purani files auto-delete)."""

    def __init__(self, log_dir):
        self.dir = log_dir
        self._day = ""
        self._events_path = ""
        self._log_path = ""
        self._ef = None      # events file handle (open rehta hai)
        self._lf = None      # plain log file handle
        try:
            os.makedirs(self.dir, exist_ok=True)
        except Exception:  # noqa: BLE001
            pass
        self._rotate()

    def _rotate(self):
        day = time.strftime("%Y%m%d")
        if day == self._day:
            return
        self._close_handles()
        self._day = day
        self._events_path = os.path.join(self.dir, "events-%s.jsonl" % day)
        self._log_path = os.path.join(self.dir, "phonemirror-%s.log" % day)
        self.enforce_retention()

    def _close_handles(self):
        for attr in ("_ef", "_lf"):
            f = getattr(self, attr)
            if f is not None:
                try:
                    f.flush()
                    f.close()
                except Exception:  # noqa: BLE001
                    pass
                setattr(self, attr, None)

    def close(self):
        self._close_handles()

    def _handle(self, events):
        attr = "_ef" if events else "_lf"
        f = getattr(self, attr)
        if f is None or f.closed:
            path = self._events_path if events else self._log_path
            f = open(path, "a", encoding="utf-8", buffering=1)  # line-buffered
            setattr(self, attr, f)
        return f

    def write_event(self, ev):
        try:
            self._rotate()
            self._handle(True).write(json.dumps(ev, ensure_ascii=False) + "\n")
        except Exception:  # noqa: BLE001
            pass

    def write_line(self, serial, text):
        try:
            self._rotate()
            ts = datetime.datetime.now().strftime("%Y-%m-%d %H:%M:%S")
            prefix = ("[%s] " % serial) if serial else ""
            self._handle(False).write("%s  %s%s\n" % (ts, prefix, text))
        except Exception:  # noqa: BLE001
            pass

    def files(self):
        try:
            return [os.path.join(self.dir, n) for n in sorted(os.listdir(self.dir))
                    if n.startswith("events-") or n.startswith("phonemirror-")]
        except Exception:  # noqa: BLE001
            return []

    def clear(self):
        self._close_handles()
        for p in self.files():
            try:
                os.remove(p)
            except Exception:  # noqa: BLE001
                pass

    def enforce_retention(self):
        cutoff = time.time() - LOG_RETENTION_DAYS * 86400
        for p in self.files():
            try:
                if os.path.getmtime(p) < cutoff:
                    os.remove(p)
            except Exception:  # noqa: BLE001
                pass


def default_base_dir():
    """Exports + logs ki base folder."""
    home = os.path.expanduser("~")
    docs = os.path.join(home, "Documents")
    return os.path.join(docs, "PhoneMirror") if os.path.isdir(docs) else os.path.join(home, "PhoneMirror")


def crash_reports_dir(base=None):
    """Crash reports ki folder (local artifact - user attach kar sakta hai)."""
    d = os.path.join(base or default_base_dir(), "crash-reports")
    try:
        os.makedirs(d, exist_ok=True)
    except Exception:  # noqa: BLE001
        pass
    return d


def write_crash_report(exc_type, exc_value, tb_text, context=None, base=None):
    """Crash ka poora report file likho (traceback + context). Returns path or ''.
    Note: file local rehti hai (raw traceback), events me sirf hash jaata hai."""
    name = "crash-%s.txt" % time.strftime("%Y%m%d_%H%M%S")
    path = os.path.join(crash_reports_dir(base), name)
    lines = [
        "PhoneMirror crash report",
        "time: %s" % utc_now(),
        "app: %s v%s (%s)" % (APP_TITLE, APP_VERSION, APP_ID),
        "os: %s %s (%s)" % (platform.system(), platform.release(), platform.machine()),
        "python: %s" % sys.version.split()[0],
        "exception: %s" % exc_type,
        "message: %s" % exc_value,
    ]
    for k, v in (context or {}).items():
        lines.append("%s: %s" % (k, v))
    lines.append("")
    lines.append("--- traceback ---")
    lines.append((tb_text or "").rstrip())
    lines.append("")
    try:
        with open(path, "w", encoding="utf-8") as f:
            f.write("\n".join(lines))
        return path
    except Exception:  # noqa: BLE001
        return ""


# --------------------------------------------------------------------------- #
#  Executable / device discovery
# --------------------------------------------------------------------------- #
def find_executable(name):
    """PATH aur common install locations me exe dhoondo."""
    found = shutil.which(name)
    if found:
        return found
    candidates = []
    if IS_WINDOWS:
        local = os.environ.get("LOCALAPPDATA", "")
        profile = os.environ.get("USERPROFILE", "")
        candidates += [
            os.path.join(local, "Microsoft", "WinGet", "Links", name + ".exe"),
            os.path.join(profile, "scoop", "shims", name + ".exe"),
            os.path.join(profile, "chocolatey", "bin", name + ".exe"),
            r"C:\scrcpy\%s.exe" % name,
            r"C:\Program Files\scrcpy\%s.exe" % name,
            r"C:\Program Files (x86)\scrcpy\%s.exe" % name,
        ]
    for c in candidates:
        if c and os.path.isfile(c):
            return c
    return None


def run_adb(adb_path, args, timeout=20):
    """adb command chalao. Returns (ok, output_text)."""
    try:
        proc = subprocess.run(
            [adb_path] + args, capture_output=True, text=True,
            timeout=timeout, creationflags=CREATE_NO_WINDOW,
        )
    except Exception as exc:  # noqa: BLE001
        return False, str(exc)
    out = ((proc.stdout or "") + (proc.stderr or "")).strip()
    return proc.returncode == 0, out


def list_devices(adb_path):
    """adb devices -l se device list nikalo. Returns (list, error)."""
    try:
        proc = subprocess.run(
            [adb_path, "devices", "-l"],
            capture_output=True, text=True, timeout=15,
            creationflags=CREATE_NO_WINDOW,
        )
    except FileNotFoundError:
        return [], "adb nahi mila."
    except Exception as exc:  # noqa: BLE001
        return [], str(exc)

    if proc.returncode != 0:
        return [], (proc.stderr or "adb devices fail hua.").strip()

    devices = []
    for line in proc.stdout.splitlines():
        line = line.strip()
        if not line or line.startswith("List of devices"):
            continue
        parts = line.split()
        if len(parts) < 2:
            continue
        serial, state = parts[0], parts[1]
        model = ""
        for token in parts[2:]:
            if token.startswith("model:"):
                model = token.split(":", 1)[1].replace("_", " ")
        label = serial
        if model:
            label = "%s  (%s)" % (model, serial)
        if state != "device":
            label += "   [%s]" % state
        devices.append({"serial": serial, "state": state, "model": model, "label": label})
    return devices, None


def detect_phone_ip(adb_path, serial):
    """USB se jude phone ka Wi-Fi IP nikalo. Returns ip string or None."""
    base = ["-s", serial] if serial else []
    ok, out = run_adb(adb_path, base + ["shell", "ip", "route"])
    if ok:
        m = re.search(r"src\s+(\d+\.\d+\.\d+\.\d+)", out)
        if m:
            return m.group(1)
    ok2, out2 = run_adb(adb_path, base + ["shell", "ip", "-f", "inet", "addr", "show", "wlan0"])
    if ok2:
        m2 = re.search(r"inet\s+(\d+\.\d+\.\d+\.\d+)", out2)
        if m2:
            return m2.group(1)
    return None


# --------------------------------------------------------------------------- #
#  Presets
# --------------------------------------------------------------------------- #
PRESETS = {
    "Low Latency": {
        "max_size": 1024, "bitrate": "8M", "max_fps": 60, "codec": "h264",
        "audio": True, "audio_codec": "opus", "audio_bitrate": "128K", "audio_buffer": 25,
        "desc": "Sabse tez. 1024px, 8 Mbps, 60fps, audio 128K. Gaming / interactive.",
    },
    "Low Phone Load": {
        "max_size": 800, "bitrate": "4M", "max_fps": 30, "codec": "h264",
        "audio": True, "audio_codec": "opus", "audio_bitrate": "96K", "audio_buffer": 50,
        "desc": "Phone par sabse kam load + kam heat. 800px, 4 Mbps, 30fps.",
    },
    "Balanced": {
        "max_size": 1280, "bitrate": "8M", "max_fps": 60, "codec": "h264",
        "audio": True, "audio_codec": "opus", "audio_bitrate": "128K", "audio_buffer": 50,
        "desc": "Rozana kaam. 1280px, 8 Mbps, 60fps, clear audio.",
    },
    "High Quality": {
        "max_size": 1920, "bitrate": "16M", "max_fps": 60, "codec": "h265",
        "audio": True, "audio_codec": "opus", "audio_bitrate": "256K", "audio_buffer": 100,
        "desc": "Sharp picture + best audio. 1920px, 16 Mbps, 256K audio.",
    },
}


def build_command(scrcpy_path, serial, s):
    """Settings dict se scrcpy command line banao."""
    cmd = [scrcpy_path]
    if serial:
        cmd += ["--serial", serial]
        cmd += ["--window-title", "PhoneMirror - %s" % serial]

    src = s.get("video_source", "display")
    if src == "camera":
        cmd += ["--video-source", "camera"]
        cmd += ["--camera-facing", s.get("camera_facing", "back")]
        cmd += ["--camera-size", s.get("camera_size", "1280x720")]
        cmd += ["--camera-fps", str(s.get("camera_fps", 30))]
    else:
        cmd += ["--max-size", str(s["max_size"])]

    cmd += ["--video-bit-rate", s["bitrate"]]
    cmd += ["--max-fps", str(s["max_fps"])]
    cmd += ["--video-codec", s["codec"]]
    cmd += ["--video-buffer", "0", "--display-buffer", "0"]
    cmd += ["--render-driver", "opengl"]

    if s.get("audio"):
        cmd += ["--audio-codec", s.get("audio_codec", "opus")]
        cmd += ["--audio-bit-rate", s.get("audio_bitrate", "128K")]
        cmd += ["--audio-buffer", str(s.get("audio_buffer", 50))]
    else:
        cmd += ["--no-audio"]

    if s.get("record") and s.get("record_path"):
        cmd += ["--record", s["record_path"]]

    if s.get("no_control"):
        cmd += ["--no-control"]
    else:
        kb = s.get("keyboard_mode", "default")
        if kb and kb != "default":
            cmd += ["--keyboard", kb]

    if s.get("turn_screen_off"):
        cmd += ["--turn-screen-off"]
    if s.get("stay_awake"):
        cmd += ["--stay-awake"]
    if s.get("always_on_top"):
        cmd += ["--always-on-top"]
    if s.get("print_fps"):
        cmd += ["--print-fps"]
    return cmd


# --------------------------------------------------------------------------- #
#  Diagnostics report rendering
# --------------------------------------------------------------------------- #
def _event_lines(rep, limit=200):
    evs = rep.get("events", [])
    out = []
    for e in evs[-limit:]:
        out.append("  %s  %-5s  %-22s  %s" % (
            e.get("timestamp", ""), e.get("log_level", ""),
            e.get("event_name", ""), e.get("screen_or_route", "")))
    return out


def render_txt(rep):
    L = []
    L.append("PhoneMirror diagnostics report")
    L.append("Generated: %s" % rep["generated_at"])
    L.append("App: %s v%s" % (rep["app"]["name"], rep["app"]["version"]))
    L.append("Diagnostics collection: %s" % ("ON" if rep["diagnostics_enabled"] else "OFF"))
    L.append("Log schema: %s" % rep.get("logging", {}).get("schema_version", ""))
    L.append("")
    L.append("== Environment ==")
    for k, v in rep["environment"].items():
        L.append("  %-10s %s" % (k + ":", v))
    L.append("")
    L.append("== Tools ==")
    for k, v in rep["tools"].items():
        L.append("  %-16s %s" % (k + ":", v))
    L.append("")
    L.append("== Devices (%d) ==" % len(rep["devices"]))
    for d in rep["devices"]:
        L.append("  %s   [%s]   %s" % (d["serial"], d["state"], d["model"]))
    L.append("")
    L.append("== Settings ==")
    for k, v in rep["settings"].items():
        L.append("  %-18s %s" % (k + ":", v))
    L.append("")
    L.append("== Analytics (metrics) ==")
    for k, v in (rep.get("analytics") or {}).items():
        L.append("  %-20s %s" % (k + ":", v))
    L.append("")
    cr = rep.get("crash_reports") or []
    L.append("== Crash reports (%d) ==" % len(cr))
    for n in cr:
        L.append("  %s" % n)
    if not cr:
        L.append("  (koi crash report nahi)")
    L.append("")
    L.append("== Structured events (%d) ==" % len(rep.get("events", [])))
    if not rep.get("events"):
        L.append("  (kuch nahi - diagnostics OFF ho sakta hai)")
    L.extend(_event_lines(rep))
    L.append("")
    L.append("== Log (%d lines) ==" % len(rep["log"]))
    for e in rep["log"]:
        pref = ("[%s] " % e["serial"]) if e.get("serial") else ""
        L.append("  %s  %s%s" % (e["time"], pref, e["text"]))
    return "\n".join(L)


def render_md(rep):
    L = []
    L.append("# PhoneMirror diagnostics report")
    L.append("")
    L.append("- **Generated:** %s" % rep["generated_at"])
    L.append("- **App:** %s v%s" % (rep["app"]["name"], rep["app"]["version"]))
    L.append("- **Diagnostics:** %s" % ("ON" if rep["diagnostics_enabled"] else "OFF"))
    L.append("- **Log schema:** %s" % rep.get("logging", {}).get("schema_version", ""))
    L.append("")
    L.append("## Environment")
    L.append("")
    L.append("| Key | Value |")
    L.append("|---|---|")
    for k, v in rep["environment"].items():
        L.append("| %s | %s |" % (k, v))
    L.append("")
    L.append("## Tools")
    L.append("")
    L.append("| Key | Value |")
    L.append("|---|---|")
    for k, v in rep["tools"].items():
        L.append("| %s | %s |" % (k, str(v).replace("|", "/")))
    L.append("")
    L.append("## Devices (%d)" % len(rep["devices"]))
    L.append("")
    L.append("| Serial | State | Model |")
    L.append("|---|---|---|")
    for d in rep["devices"]:
        L.append("| %s | %s | %s |" % (d["serial"], d["state"], d["model"]))
    L.append("")
    L.append("## Settings")
    L.append("")
    L.append("| Setting | Value |")
    L.append("|---|---|")
    for k, v in rep["settings"].items():
        L.append("| %s | %s |" % (k, v))
    L.append("")
    L.append("## Analytics (metrics)")
    L.append("")
    L.append("| Metric | Value |")
    L.append("|---|---|")
    for k, v in (rep.get("analytics") or {}).items():
        L.append("| %s | %s |" % (k, v))
    L.append("")
    cr = rep.get("crash_reports") or []
    L.append("## Crash reports (%d)" % len(cr))
    L.append("")
    if cr:
        for n in cr:
            L.append("- `%s`" % n)
    else:
        L.append("_koi crash report nahi_")
    L.append("")
    L.append("## Structured events (%d)" % len(rep.get("events", [])))
    L.append("")
    L.append("| timestamp | level | event_name | screen_or_route |")
    L.append("|---|---|---|---|")
    for e in rep.get("events", [])[-200:]:
        L.append("| %s | %s | %s | %s |" % (e.get("timestamp", ""), e.get("log_level", ""),
                                            e.get("event_name", ""), e.get("screen_or_route", "")))
    L.append("")
    L.append("## Log (%d lines)" % len(rep["log"]))
    L.append("")
    L.append("```")
    for e in rep["log"]:
        pref = ("[%s] " % e["serial"]) if e.get("serial") else ""
        L.append("%s  %s%s" % (e["time"], pref, e["text"]))
    L.append("```")
    return "\n".join(L)


def render_email_body(rep, attach_path, max_lines=40, limit=1800):
    lines = [
        "PhoneMirror diagnostics report",
        "Generated: %s" % rep["generated_at"],
        "App: %s v%s" % (rep["app"]["name"], rep["app"]["version"]),
        "Diagnostics: %s" % ("ON" if rep["diagnostics_enabled"] else "OFF"),
        "Log schema: %s" % rep.get("logging", {}).get("schema_version", ""),
        "",
        "OS: %s" % rep["environment"].get("os", ""),
        "Python: %s" % rep["environment"].get("python", ""),
        "scrcpy: %s" % rep["tools"].get("scrcpy_version", ""),
        "adb: %s" % rep["tools"].get("adb_version", ""),
        "Devices: %d" % len(rep["devices"]),
        "Structured events: %d" % len(rep.get("events", [])),
        "",
        "--- Log (last %d lines) ---" % max_lines,
    ]
    for e in rep["log"][-max_lines:]:
        pref = ("[%s] " % e["serial"]) if e.get("serial") else ""
        lines.append("%s %s%s" % (e["time"], pref, e["text"]))
    lines.append("")
    lines.append("(Poora report file attach kiya gaya hai: %s)" % os.path.basename(attach_path))
    body = "\n".join(lines)
    if len(body) > limit:
        body = body[:limit] + "\n...(truncated - poora report file me hai)"
    return body


def write_xlsx(rep, path):
    """xlsx likho. openpyxl chahiye. Raises ImportError agar nahi hai."""
    from openpyxl import Workbook  # noqa: WPS433
    wb = Workbook()
    ws = wb.active
    ws.title = "Summary"
    ws.append(["Field", "Value"])
    ws.append(["Generated", rep["generated_at"]])
    ws.append(["App", "%s v%s" % (rep["app"]["name"], rep["app"]["version"])])
    ws.append(["Diagnostics", "ON" if rep["diagnostics_enabled"] else "OFF"])
    ws.append(["Log schema", rep.get("logging", {}).get("schema_version", "")])
    for k, v in rep["environment"].items():
        ws.append(["env." + k, str(v)])
    for k, v in rep["tools"].items():
        ws.append(["tool." + k, str(v)])

    dev = wb.create_sheet("Devices")
    dev.append(["Serial", "State", "Model"])
    for d in rep["devices"]:
        dev.append([d["serial"], d["state"], d["model"]])

    st = wb.create_sheet("Settings")
    st.append(["Setting", "Value"])
    for k, v in rep["settings"].items():
        st.append([k, str(v)])

    an = wb.create_sheet("Analytics")
    an.append(["Metric", "Value"])
    for k, v in (rep.get("analytics") or {}).items():
        an.append([k, v])

    crs = wb.create_sheet("CrashReports")
    crs.append(["Crash report file"])
    for n in (rep.get("crash_reports") or []):
        crs.append([n])

    ev = wb.create_sheet("Events")
    ev.append(["timestamp", "log_level", "event_name", "screen_or_route", "attempt_id", "session_id"])
    for e in rep.get("events", []):
        ev.append([e.get("timestamp", ""), e.get("log_level", ""), e.get("event_name", ""),
                   e.get("screen_or_route", ""), e.get("attempt_id", ""), e.get("session_id", "")])

    lg = wb.create_sheet("Log")
    lg.append(["Time", "Serial", "Message"])
    for e in rep["log"]:
        lg.append([e["time"], e.get("serial", ""), e["text"]])

    wb.save(path)


def write_csv(rep, path):
    with open(path, "w", newline="", encoding="utf-8") as f:
        w = csv.writer(f)
        w.writerow(["time", "serial", "message"])
        for e in rep["log"]:
            w.writerow([e["time"], e.get("serial", ""), e["text"]])


def write_jsonl(rep, path):
    with open(path, "w", encoding="utf-8") as f:
        f.write("\n".join(json.dumps(e, ensure_ascii=False) for e in rep.get("events", [])))


def write_zip(rep, path, logs_dir=None, crash_dir=None):
    """Poora bundle ek zip me: report + events + device logs + crash reports."""
    import zipfile  # noqa: WPS433
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as z:
        z.writestr("report.md", render_md(rep))
        z.writestr("report.txt", render_txt(rep))
        z.writestr("report.json", json.dumps(rep, indent=2, ensure_ascii=False))
        z.writestr("events.jsonl",
                   "\n".join(json.dumps(e, ensure_ascii=False) for e in rep.get("events", [])))
        if logs_dir and os.path.isdir(logs_dir):
            for name in sorted(os.listdir(logs_dir)):
                if name.startswith("events-") or name.startswith("phonemirror-"):
                    try:
                        z.write(os.path.join(logs_dir, name), arcname="device_logs/" + name)
                    except Exception:  # noqa: BLE001
                        pass
        if crash_dir and os.path.isdir(crash_dir):
            for name in sorted(os.listdir(crash_dir)):
                if name.startswith("crash-"):
                    try:
                        z.write(os.path.join(crash_dir, name), arcname="crash_reports/" + name)
                    except Exception:  # noqa: BLE001
                        pass


# --------------------------------------------------------------------------- #
#  Session
# --------------------------------------------------------------------------- #
class Session:
    """Ek device ka mirroring session."""

    def __init__(self, serial, label=""):
        self.serial = serial
        self.label = label or serial
        self.proc = None
        self.start_time = 0.0
        self.reconnect_attempts = 0
        self.user_stopped = False
        self.running_intent = False

    def active(self):
        return self.running_intent or self.proc is not None


# --------------------------------------------------------------------------- #
#  Diagnostics window
# --------------------------------------------------------------------------- #
class DiagnosticsWindow(tk.Toplevel):
    def __init__(self, app):
        super().__init__(app.root)
        self.app = app
        self.title("Diagnostics  -  PhoneMirror")
        self.minsize(780, 620)
        self._last_path = None

        pad = {"padx": 8, "pady": 4}

        top = ttk.LabelFrame(self, text="Settings")
        top.pack(fill="x", **pad)
        ttk.Checkbutton(top, text="Collect diagnostics (logs + technical data)  -  default ON",
                        variable=app.diag_enabled, command=self._toggle).grid(
            row=0, column=0, sticky="w", padx=6, pady=6)
        ttk.Label(top, text="Off karne par app koi data collect nahi karegi (export sirf summary dega). "
                            "Structured events ProGallery Logging Spec v1.0 ke hisaab se hain.",
                  foreground="#666").grid(row=1, column=0, sticky="w", padx=8, pady=(0, 6))

        self.preview = tk.Text(self, height=18, wrap="none", bg="#111", fg="#d0d0d0",
                               insertbackground="#d0d0d0", font=("Consolas", 9))
        self.preview.pack(fill="both", expand=True, padx=8, pady=4)
        self.preview.configure(state="disabled")

        bar = ttk.Frame(self)
        bar.pack(fill="x", **pad)
        ttk.Button(bar, text="Refresh", command=self.refresh).pack(side="left", padx=3)
        ttk.Button(bar, text="Clean", command=self.clean).pack(side="left", padx=3)
        ttk.Button(bar, text="Copy report", command=self.copy).pack(side="left", padx=3)

        ttk.Label(bar, text="Format:").pack(side="left", padx=(16, 4))
        self.fmt = tk.StringVar(value="txt")
        ttk.Combobox(bar, textvariable=self.fmt, state="readonly", width=8,
                     values=["txt", "md", "json", "jsonl", "csv", "xlsx", "zip"]).pack(side="left")
        ttk.Button(bar, text="Save as...", command=self.save_as).pack(side="left", padx=6)

        bar2 = ttk.Frame(self)
        bar2.pack(fill="x", **pad)
        ttk.Button(bar2, text="Share bundle", command=self.share).pack(side="left", padx=3)
        ttk.Button(bar2, text="Email (Gmail)", command=lambda: self.email(True)).pack(side="left", padx=3)
        ttk.Button(bar2, text="Email (mail app)", command=lambda: self.email(False)).pack(side="left", padx=3)
        ttk.Button(bar2, text="Open exports", command=self.open_folder).pack(side="left", padx=3)
        ttk.Button(bar2, text="Open logs", command=self.open_logs).pack(side="left", padx=3)
        ttk.Label(bar2, text="Email: %s" % DIAG_EMAIL, foreground="#666").pack(side="right", padx=6)

        self.status = ttk.Label(self, text="", anchor="w", relief="sunken")
        self.status.pack(fill="x", side="bottom")

        self.refresh()

    def _toggle(self):
        self.app.set_diag_enabled(self.app.diag_enabled.get())
        self.status.configure(text="Diagnostics %s." % ("ON" if self.app.diag_enabled.get() else "OFF"))
        self.refresh()

    def _render(self):
        rep = self.app.build_report()
        return render_txt(rep), rep

    def refresh(self):
        text, _ = self._render()
        self.preview.configure(state="normal")
        self.preview.delete("1.0", "end")
        self.preview.insert("end", text)
        self.preview.configure(state="disabled")
        self.status.configure(text="Report refreshed - %d log lines, %d events."
                              % (len(self.app.diag), len(self.app.logger.events)))

    def clean(self):
        if not messagebox.askyesno("Clean", "Diagnostics log aur collected data clear kar dun?"):
            return
        self.app.clean_diagnostics()
        self.refresh()
        self.status.configure(text="Diagnostics data clear ho gaya.")

    def copy(self):
        text, _ = self._render()
        self.app.root.clipboard_clear()
        self.app.root.clipboard_append(text)
        self.status.configure(text="Poora report clipboard me copy ho gaya.")

    def save_as(self):
        fmt = self.fmt.get()
        ext = "." + fmt
        path = filedialog.asksaveasfilename(
            title="Diagnostics report save karo", defaultextension=ext,
            initialfile="PhoneMirror_diagnostics_%s%s" % (time.strftime("%Y%m%d_%H%M%S"), ext),
            filetypes=[(fmt.upper(), "*" + ext), ("All", "*.*")])
        if not path:
            return
        _, rep = self._render()
        try:
            self.app.save_report(rep, path, fmt)
            self._last_path = path
            self.status.configure(text="Saved: %s" % path)
            self.app._log(None, "[i] Diagnostics saved: %s" % path)
        except Exception as exc:  # noqa: BLE001
            self.app.toast("Save fail: %s" % exc, "error")

    def email(self, gmail):
        try:
            path = self.app.email_report(gmail)
            self._last_path = path
            self.status.configure(text="Email compose khula. Report: %s" % path)
        except Exception as exc:  # noqa: BLE001
            self.app.toast("Email fail: %s" % exc, "error")

    def open_folder(self):
        d = self.app.diagnostics_dir()
        try:
            if IS_WINDOWS:
                os.startfile(d)  # noqa: S606
            else:
                webbrowser.open("file://" + d)
        except Exception:  # noqa: BLE001
            messagebox.showinfo("Folder", d)

    def open_logs(self):
        d = self.app.logs_dir()
        try:
            if IS_WINDOWS:
                os.startfile(d)  # noqa: S606
            else:
                webbrowser.open("file://" + d)
        except Exception:  # noqa: BLE001
            messagebox.showinfo("Logs folder", d)

    def share(self):
        try:
            path = self.app.share_bundle()
            self._last_path = path
            self.status.configure(text="Share bundle: %s (path clipboard me)" % path)
        except Exception as exc:  # noqa: BLE001
            self.app.toast("Share fail: %s" % exc, "error")


# --------------------------------------------------------------------------- #
#  Main app
# --------------------------------------------------------------------------- #
class PhoneMirrorApp:
    def __init__(self, root):
        self.root = root
        self.root.title(APP_TITLE)
        self.root.minsize(880, 820)

        cfg = load_config()
        salt = cfg.get("salt") or _rand_hex(32)
        if cfg.get("salt") != salt:
            cfg["salt"] = salt
            save_config(cfg)
        self._salt = salt

        self.log_q = queue.Queue()
        self.devices = []
        self.sessions = {}
        self.diag = []                      # collected human-readable log entries
        self._diag_win = None
        self._launch_attempt = None
        self._launch_multi = False
        self._refreshing = False
        self._probe = {"scrcpy_version": "", "adb_version": ""}
        self._want_serials = []
        self._toast_job = None
        self._last_screen = "Main"
        self._last_action = "app_start"
        # analytics counters (metrics; raw logs se alag)
        self.stats = {
            "mirrors_started": 0, "mirrors_succeeded": 0, "mirrors_failed": 0,
            "devices_seen": 0, "reconnects": 0, "exports": 0, "emails": 0,
            "shares": 0, "errors": 0, "crashes": 0,
        }
        self.session_id = "sess_" + _rand_hex(12)
        self.logger = StructuredLogger(self.session_id)
        self.filelog = FileLogWriter(os.path.join(default_base_dir(), "logs"))
        self.scrcpy_path = tk.StringVar(value=find_executable("scrcpy") or "")
        self.adb_path = tk.StringVar(value=find_executable("adb") or "")

        self.diag_enabled = tk.BooleanVar(value=bool(cfg.get("diagnostics_enabled", True)))

        self.connection_mode = tk.StringVar(value="usb")
        self.wifi_ip = tk.StringVar(value="")
        self.wifi_port = tk.IntVar(value=5555)

        self.video_source = tk.StringVar(value="display")
        self.max_size = tk.IntVar(value=1024)
        self.bitrate = tk.StringVar(value="8M")
        self.max_fps = tk.IntVar(value=60)
        self.codec = tk.StringVar(value="h264")
        self.camera_facing = tk.StringVar(value="back")
        self.camera_size = tk.StringVar(value="1280x720")
        self.camera_fps = tk.IntVar(value=30)

        self.audio = tk.BooleanVar(value=True)
        self.audio_codec = tk.StringVar(value="opus")
        self.audio_bitrate = tk.StringVar(value="128K")
        self.audio_buffer = tk.IntVar(value=50)

        self.record = tk.BooleanVar(value=False)
        self.record_path = tk.StringVar(value="")

        self.no_control = tk.BooleanVar(value=False)
        self.keyboard_mode = tk.StringVar(value="default")

        self.turn_screen_off = tk.BooleanVar(value=False)
        self.stay_awake = tk.BooleanVar(value=True)
        self.always_on_top = tk.BooleanVar(value=False)
        self.print_fps = tk.BooleanVar(value=False)

        self.auto_reconnect = tk.BooleanVar(value=True)
        self.auto_refresh = tk.BooleanVar(value=True)
        self.reconnect_delay = tk.IntVar(value=3)
        self.reconnect_max = tk.IntVar(value=0)

        self._build_ui()
        self._apply_saved(cfg.get("settings", {}))
        self._poll_log()
        self._ev("session_start", "info", "Main", props={"start_type": "cold"})
        self._ev("environment_snapshot", "info", "Main",
                 props={"os": _env_os(), "python": sys.version.split()[0],
                        "machine": _device_model()})
        self.refresh_probe()
        self.refresh_devices()
        self.root.after(5000, self._auto_refresh_tick)

    # ---------------- structured events ----------------
    def _ev(self, event_name, log_level, screen_or_route, attempt_id=None, props=None,
            category="", message=""):
        """Structured event log (sirf diagnostics ON hone par) + device file me save.
        category event_name se infer hoti hai (ya explicit override)."""
        if not self.diag_enabled.get():
            return None
        if log_level in ("error", "fatal"):
            self.stats["errors"] += 1
        ev = self.logger.log(event_name, log_level, screen_or_route,
                             attempt_id=attempt_id, props=props,
                             category=category_for(event_name, category), message=message)
        if ev is not None:
            try:
                self.filelog.write_event(ev)
            except Exception:  # noqa: BLE001
                pass
        return ev

    def _dev_pseudonym(self, serial):
        return salted_hash(serial, self._salt)

    def toast(self, msg, level="info", ms=4000):
        """Non-modal notification - status bar ke upar chhota toast (flow block nahi hota)."""
        colors = {"info": "#2b6cb0", "warn": "#b7791f", "error": "#c53030", "success": "#2f855a"}
        try:
            self._toast_lbl.configure(text=msg, background=colors.get(level, "#2b6cb0"))
            self._toast_lbl.place(relx=0.0, rely=1.0, anchor="sw", relwidth=1.0)
            if getattr(self, "_toast_job", None):
                try:
                    self.root.after_cancel(self._toast_job)
                except Exception:  # noqa: BLE001
                    pass
            self._toast_job = self.root.after(ms, self._hide_toast)
        except Exception:  # noqa: BLE001
            pass

    def _hide_toast(self):
        try:
            self._toast_lbl.place_forget()
        except Exception:  # noqa: BLE001
            pass
        self._toast_job = None

    def analytics_snapshot(self, where="manual"):
        """Analytics metrics ka point-in-time snapshot event (category: analytics)."""
        props = dict(self.stats)
        props["where"] = where
        props["session_id"] = self.session_id
        self._ev("analytics_snapshot", "info", "Main", props=props)
        return dict(self.stats)

    # ---------------- UI ----------------
    def _build_ui(self):
        pad = {"padx": 8, "pady": 3}

        paths = ttk.LabelFrame(self.root, text="1. Tools (scrcpy / adb)")
        paths.pack(fill="x", **pad)
        ttk.Label(paths, text="scrcpy:").grid(row=0, column=0, sticky="w", padx=6, pady=3)
        ttk.Entry(paths, textvariable=self.scrcpy_path, width=52).grid(row=0, column=1, sticky="we", pady=3)
        ttk.Button(paths, text="Browse", command=self._browse_scrcpy).grid(row=0, column=2, padx=6)
        ttk.Label(paths, text="adb:").grid(row=1, column=0, sticky="w", padx=6, pady=3)
        ttk.Entry(paths, textvariable=self.adb_path, width=52).grid(row=1, column=1, sticky="we", pady=3)
        ttk.Button(paths, text="Browse", command=self._browse_adb).grid(row=1, column=2, padx=6)
        paths.columnconfigure(1, weight=1)

        dev = ttk.LabelFrame(self.root, text="2. Device  (multi-select: Ctrl / Shift + click)")
        dev.pack(fill="x", **pad)
        self.device_list = tk.Listbox(dev, selectmode="extended", height=4, exportselection=False)
        self.device_list.grid(row=0, column=0, sticky="we", padx=6, pady=6)
        sb = ttk.Scrollbar(dev, orient="vertical", command=self.device_list.yview)
        sb.grid(row=0, column=1, sticky="ns", pady=6)
        self.device_list.configure(yscrollcommand=sb.set)
        ttk.Button(dev, text="Refresh", command=self.refresh_devices).grid(row=0, column=2, padx=6, sticky="n")
        ttk.Checkbutton(dev, text="Auto-refresh", variable=self.auto_refresh).grid(row=1, column=2, padx=6, sticky="e")
        self.device_status = ttk.Label(dev, text="", foreground="#666")
        self.device_status.grid(row=1, column=0, columnspan=2, sticky="w", padx=8)
        dev.columnconfigure(0, weight=1)

        conn = ttk.LabelFrame(self.root, text="3. Connection (USB / Wi-Fi)")
        conn.pack(fill="x", **pad)
        ttk.Radiobutton(conn, text="USB cable  (lowest latency)", value="usb",
                        variable=self.connection_mode).grid(row=0, column=0, columnspan=5, sticky="w", padx=6)
        ttk.Radiobutton(conn, text="Wi-Fi / TCP-IP  (no cable)", value="wifi",
                        variable=self.connection_mode).grid(row=1, column=0, columnspan=5, sticky="w", padx=6)
        ttk.Label(conn, text="Phone IP:").grid(row=2, column=0, sticky="e", padx=(16, 4), pady=4)
        ttk.Entry(conn, textvariable=self.wifi_ip, width=15).grid(row=2, column=1, sticky="w")
        ttk.Label(conn, text="Port:").grid(row=2, column=2, sticky="e", padx=(12, 4))
        ttk.Entry(conn, textvariable=self.wifi_port, width=7).grid(row=2, column=3, sticky="w")
        ttk.Button(conn, text="Auto-detect IP", command=self.detect_ip).grid(row=2, column=4, padx=8)
        ttk.Button(conn, text="Enable Wi-Fi", command=self.enable_wifi).grid(row=3, column=0, padx=6, pady=(0, 4), sticky="w")
        ttk.Button(conn, text="Connect", command=self.connect_wifi).grid(row=3, column=1, padx=6, pady=(0, 4), sticky="w")
        ttk.Button(conn, text="Disconnect", command=self.disconnect_wifi).grid(row=3, column=2, padx=6, pady=(0, 4), sticky="w")

        pre = ttk.LabelFrame(self.root, text="4. Preset")
        pre.pack(fill="x", **pad)
        for i, name in enumerate(PRESETS):
            ttk.Button(pre, text=name, command=lambda n=name: self.apply_preset(n)).grid(
                row=0, column=i, padx=6, pady=4, sticky="we")
            pre.columnconfigure(i, weight=1)
        self.preset_desc = ttk.Label(pre, text=PRESETS["Low Latency"]["desc"],
                                     wraplength=800, foreground="#444")
        self.preset_desc.grid(row=1, column=0, columnspan=len(PRESETS), sticky="w", padx=8, pady=(0, 4))

        tun = ttk.LabelFrame(self.root, text="5. Video")
        tun.pack(fill="x", **pad)
        srcrow = ttk.Frame(tun)
        srcrow.grid(row=0, column=0, columnspan=4, sticky="w", padx=6, pady=(4, 0))
        ttk.Label(srcrow, text="Source:").pack(side="left")
        ttk.Radiobutton(srcrow, text="Display", value="display", variable=self.video_source,
                        command=self._on_source).pack(side="left", padx=(8, 0))
        ttk.Radiobutton(srcrow, text="Camera", value="camera", variable=self.video_source,
                        command=self._on_source).pack(side="left", padx=(8, 0))
        ttk.Label(tun, text="Max size (px):").grid(row=1, column=0, sticky="w", padx=6, pady=3)
        self.size_scale = ttk.Scale(tun, from_=480, to=2560, variable=self.max_size,
                                    orient="horizontal", command=self._on_size)
        self.size_scale.grid(row=1, column=1, sticky="we", padx=6)
        self.size_lbl = ttk.Label(tun, text="1024", width=6)
        self.size_lbl.grid(row=1, column=2, sticky="w")
        ttk.Label(tun, text="Bitrate:").grid(row=2, column=0, sticky="w", padx=6, pady=3)
        ttk.Combobox(tun, textvariable=self.bitrate, state="readonly", width=10,
                     values=["2M", "4M", "6M", "8M", "12M", "16M", "24M"]).grid(row=2, column=1, sticky="w", padx=6)
        ttk.Label(tun, text="Max FPS:").grid(row=3, column=0, sticky="w", padx=6, pady=3)
        ttk.Combobox(tun, textvariable=self.max_fps, state="readonly", width=10,
                     values=[15, 30, 45, 60, 90, 120]).grid(row=3, column=1, sticky="w", padx=6)
        ttk.Label(tun, text="Codec:").grid(row=4, column=0, sticky="w", padx=6, pady=3)
        ttk.Combobox(tun, textvariable=self.codec, state="readonly", width=10,
                     values=["h264", "h265", "av1"]).grid(row=4, column=1, sticky="w", padx=6)
        cam = ttk.Frame(tun)
        cam.grid(row=1, column=3, rowspan=2, sticky="nw", padx=16)
        ttk.Label(cam, text="Camera facing:").grid(row=0, column=0, sticky="e", padx=(0, 4))
        ttk.Combobox(cam, textvariable=self.camera_facing, state="readonly", width=9,
                     values=["back", "front", "external"]).grid(row=0, column=1, sticky="w")
        ttk.Label(cam, text="Camera size:").grid(row=1, column=0, sticky="e", padx=(0, 4), pady=(4, 0))
        ttk.Combobox(cam, textvariable=self.camera_size, state="readonly", width=11,
                     values=["1920x1080", "1280x720", "640x480", "2560x1440"]).grid(row=1, column=1, sticky="w", pady=(4, 0))
        ttk.Label(cam, text="Camera FPS:").grid(row=2, column=0, sticky="e", padx=(0, 4), pady=(4, 0))
        ttk.Combobox(cam, textvariable=self.camera_fps, state="readonly", width=9,
                     values=[15, 24, 30, 60]).grid(row=2, column=1, sticky="w", pady=(4, 0))
        vopts = ttk.Frame(tun)
        vopts.grid(row=3, column=3, rowspan=2, sticky="nw", padx=16)
        ttk.Checkbutton(vopts, text="Turn phone screen off", variable=self.turn_screen_off).pack(anchor="w")
        ttk.Checkbutton(vopts, text="Stay awake", variable=self.stay_awake).pack(anchor="w")
        ttk.Checkbutton(vopts, text="Always on top", variable=self.always_on_top).pack(anchor="w")
        ttk.Checkbutton(vopts, text="Print FPS (debug)", variable=self.print_fps).pack(anchor="w")
        tun.columnconfigure(1, weight=1)

        aud = ttk.LabelFrame(self.root, text="6. Audio")
        aud.pack(fill="x", **pad)
        ttk.Checkbutton(aud, text="Audio on (clear)", variable=self.audio).grid(row=0, column=0, sticky="w", padx=6, pady=4)
        ttk.Label(aud, text="Codec:").grid(row=0, column=1, sticky="e", padx=(16, 4))
        ttk.Combobox(aud, textvariable=self.audio_codec, state="readonly", width=8,
                     values=["opus", "aac", "flac", "raw"]).grid(row=0, column=2, sticky="w")
        ttk.Label(aud, text="Bitrate:").grid(row=0, column=3, sticky="e", padx=(16, 4))
        ttk.Combobox(aud, textvariable=self.audio_bitrate, state="readonly", width=8,
                     values=["64K", "96K", "128K", "192K", "256K", "320K"]).grid(row=0, column=4, sticky="w")
        ttk.Label(aud, text="Buffer (ms):").grid(row=0, column=5, sticky="e", padx=(16, 4))
        ttk.Combobox(aud, textvariable=self.audio_buffer, state="readonly", width=6,
                     values=[0, 25, 50, 100, 150, 200]).grid(row=0, column=6, sticky="w", padx=(0, 6))

        rec = ttk.LabelFrame(self.root, text="7. Recording")
        rec.pack(fill="x", **pad)
        ttk.Checkbutton(rec, text="Record session to file", variable=self.record).grid(
            row=0, column=0, sticky="w", padx=6, pady=4)
        ttk.Entry(rec, textvariable=self.record_path, width=52).grid(row=0, column=1, sticky="we", padx=6)
        ttk.Button(rec, text="Save as...", command=self._browse_record).grid(row=0, column=2, padx=6)
        rec.columnconfigure(1, weight=1)

        rc = ttk.LabelFrame(self.root, text="8. Auto-reconnect (per device)")
        rc.pack(fill="x", **pad)
        ttk.Checkbutton(rc, text="Auto-reconnect on disconnect", variable=self.auto_reconnect).grid(
            row=0, column=0, sticky="w", padx=6, pady=4)
        ttk.Label(rc, text="Retry delay (s):").grid(row=0, column=1, sticky="e", padx=(16, 4))
        ttk.Spinbox(rc, from_=1, to=60, textvariable=self.reconnect_delay, width=5).grid(row=0, column=2, sticky="w")
        ttk.Label(rc, text="Max attempts (0 = unlimited):").grid(row=0, column=3, sticky="e", padx=(16, 4))
        ttk.Spinbox(rc, from_=0, to=1000, textvariable=self.reconnect_max, width=6).grid(row=0, column=4, sticky="w", padx=(0, 6))

        ctl = ttk.LabelFrame(self.root, text="9. Input control (PC se phone chalao)")
        ctl.pack(fill="x", **pad)
        ttk.Checkbutton(ctl, text="No control (sirf dekho)", variable=self.no_control).grid(
            row=0, column=0, sticky="w", padx=6, pady=4)
        ttk.Label(ctl, text="Keyboard mode:").grid(row=0, column=1, sticky="e", padx=(16, 4))
        ttk.Combobox(ctl, textvariable=self.keyboard_mode, state="readonly", width=10,
                     values=["default", "uhid", "sdk", "disabled"]).grid(row=0, column=2, sticky="w")

        act = ttk.Frame(self.root)
        act.pack(fill="x", **pad)
        self.launch_btn = ttk.Button(act, text="\u25b6  Start Mirroring", command=self.launch)
        self.launch_btn.pack(side="left", padx=4)
        self.stop_btn = ttk.Button(act, text="\u25a0  Stop All", command=self.stop, state="disabled")
        self.stop_btn.pack(side="left", padx=4)
        ttk.Button(act, text="Copy command(s)", command=self.copy_command).pack(side="left", padx=4)
        ttk.Button(act, text="\u2699  Diagnostics", command=self.open_diagnostics).pack(side="right", padx=4)

        logf = ttk.LabelFrame(self.root, text="Log")
        logf.pack(fill="both", expand=True, **pad)
        self.log = tk.Text(logf, height=8, wrap="word", bg="#111", fg="#d0d0d0",
                           insertbackground="#d0d0d0", font=("Consolas", 9))
        self.log.pack(side="left", fill="both", expand=True, padx=(6, 0), pady=6)
        lsb = ttk.Scrollbar(logf, command=self.log.yview)
        lsb.pack(side="right", fill="y", pady=6, padx=(0, 6))
        self.log.configure(yscrollcommand=lsb.set, state="disabled")

        self.status = ttk.Label(self.root, text="Ready.", anchor="w", relief="sunken")
        self.status.pack(fill="x", side="bottom")
        # non-modal toast (status bar ke upar)
        self._toast_lbl = tk.Label(self.root, text="", anchor="w", padx=10, pady=6,
                                   background="#2b6cb0", foreground="white")

    # ---------------- small helpers ----------------
    def _on_size(self, _val):
        self.size_lbl.configure(text=str(int(float(self.max_size.get()))))

    def _on_source(self):
        cam = (self.video_source.get() == "camera")
        self.size_scale.configure(state="disabled" if cam else "normal")

    def _browse_scrcpy(self):
        p = filedialog.askopenfilename(title="scrcpy.exe chuno",
                                       filetypes=[("scrcpy", "scrcpy*"), ("All", "*.*")])
        if p:
            self.scrcpy_path.set(p)

    def _browse_adb(self):
        p = filedialog.askopenfilename(title="adb.exe chuno",
                                       filetypes=[("adb", "adb*"), ("All", "*.*")])
        if p:
            self.adb_path.set(p)

    def _browse_record(self):
        p = filedialog.asksaveasfilename(title="Recording kahan save karein",
                                         defaultextension=".mp4",
                                         filetypes=[("MP4 video", "*.mp4"), ("MKV video", "*.mkv")])
        if p:
            self.record_path.set(p)
            self.record.set(True)

    def _log(self, serial, text):
        prefix = ("[%s] " % serial) if serial else ""
        self.log.configure(state="normal")
        self.log.insert("end", prefix + text + "\n")
        self.log.see("end")
        self.log.configure(state="disabled")
        if self.diag_enabled.get():
            safe = redact(text)
            self.diag.append({
                "time": datetime.datetime.now().strftime("%H:%M:%S"),
                "serial": serial or "",
                "text": safe,
            })
            if len(self.diag) > 5000:
                del self.diag[:1000]
            try:
                self.filelog.write_line(serial, safe)
            except Exception:  # noqa: BLE001
                pass

    def _adb(self):
        adb = self.adb_path.get().strip()
        if not adb or not os.path.exists(adb):
            adb = find_executable("adb")
            if adb:
                self.adb_path.set(adb)
        return adb

    def _run_adb_logged(self, adb, args, screen="Main"):
        t0 = time.time()
        ok, out = run_adb(adb, args)
        self._ev("adb_command", "debug", screen, props={
            "command": "adb " + " ".join(args),
            "outcome": "success" if ok else "failed",
            "latency_ms": int((time.time() - t0) * 1000),
        })
        return ok, out

    def _run_bg(self, fn, on_done):
        """fn() ko background thread me chalao (UI freeze na ho).
        Result main thread par on_done(ok, res) se aata hai (queue ke through)."""
        def work():
            try:
                res = fn()
                self.log_q.put(("call", lambda: on_done(True, res)))
            except Exception as exc:  # noqa: BLE001
                msg = str(exc)
                self.log_q.put(("call", lambda: on_done(False, msg)))
        threading.Thread(target=work, daemon=True).start()

    def _adb_bg(self, args, screen, on_done):
        """adb command background me + adb_command event log."""
        adb = self._adb()
        if not adb:
            on_done(False, {"ok": False, "out": "adb nahi mila", "ms": 0})
            return

        def work():
            t0 = time.time()
            ok, out = run_adb(adb, args)
            return {"ok": ok, "out": out, "ms": int((time.time() - t0) * 1000)}

        def done(ok, res):
            if ok:
                self._ev("adb_command", "debug", screen, props={
                    "command": "adb " + " ".join(args),
                    "outcome": "success" if res["ok"] else "failed",
                    "latency_ms": res["ms"],
                })
                self._ev("performance_marker", "info", screen,
                         props={"flow": "adb_command", "duration_ms": res["ms"]})
            on_done(ok, res)

        self._run_bg(work, done)

    # ---------------- diagnostics ----------------
    def set_diag_enabled(self, value):
        value = bool(value)
        if not value:
            # off karne se pehle audit entry likho (warna collect band ho jata hai)
            self._ev("diagnostics_toggled", "warn", "Diagnostics", props={"enabled": False})
        self.diag_enabled.set(value)
        cfg = load_config()
        cfg["diagnostics_enabled"] = value
        save_config(cfg)
        if value:
            self._ev("diagnostics_toggled", "info", "Diagnostics", props={"enabled": True})

    def clean_diagnostics(self):
        self._ev("diagnostics_cleaned", "warn", "Diagnostics",
                 props={"log_lines": len(self.diag), "events": len(self.logger.events)})
        self.diag = []
        self.logger.events = []
        self.log.configure(state="normal")
        self.log.delete("1.0", "end")
        self.log.configure(state="disabled")
        try:
            self.filelog.clear()
        except Exception:  # noqa: BLE001
            pass

    def diagnostics_dir(self):
        d = default_base_dir()
        try:
            os.makedirs(d, exist_ok=True)
        except Exception:  # noqa: BLE001
            d = os.path.expanduser("~")
        return d

    def logs_dir(self):
        d = os.path.join(self.diagnostics_dir(), "logs")
        try:
            os.makedirs(d, exist_ok=True)
        except Exception:  # noqa: BLE001
            pass
        return d

    def crash_reports_dir(self):
        return crash_reports_dir(self.diagnostics_dir())

    def refresh_probe(self):
        """scrcpy/adb version background me probe karo (UI block na ho)."""
        def work():
            scrcpy = self._resolve_scrcpy()
            sv = ""
            if scrcpy:
                try:
                    out = subprocess.run([scrcpy, "--version"], capture_output=True, text=True,
                                         timeout=10, creationflags=CREATE_NO_WINDOW).stdout.strip()
                    sv = out.splitlines()[0] if out else ""
                except Exception as exc:  # noqa: BLE001
                    sv = "error: %s" % exc
            adb = self._adb()
            av = ""
            if adb:
                ok, out = run_adb(adb, ["version"])
                av = out.splitlines()[0] if out else ""
            return {"scrcpy_version": sv, "adb_version": av}

        def done(ok, res):
            if ok:
                self._probe.update(res)
                self._ev("tool_probe", "debug", "Main",
                         props={"outcome": "success",
                                "scrcpy_version": res.get("scrcpy_version", ""),
                                "adb_version": res.get("adb_version", "")})

        self._run_bg(work, done)

    def build_report(self):
        """Report banao - koi blocking subprocess call nahi (cached probe + device list)."""
        scrcpy = self._resolve_scrcpy()
        adb = self._adb()
        devices = [{"serial": d["serial"], "state": d["state"], "model": d["model"]}
                   for d in self.devices]

        return {
            "generated_at": datetime.datetime.now().astimezone().isoformat(timespec="seconds"),
            "app": {"name": "PhoneMirror", "version": APP_VERSION},
            "diagnostics_enabled": bool(self.diag_enabled.get()),
            "logging": {
                "schema_version": LOG_SCHEMA_VERSION,
                "session_id": self.session_id,
                "retention_days": LOG_RETENTION_DAYS,
                "sample_info_rate": LOG_SAMPLE_INFO_RATE,
                "release_channel": LOG_RELEASE_CHANNEL,
            },
            "environment": {
                # platform.platform() internally subprocess use karta hai -> cheaper composite
                "os": "%s %s (%s)" % (platform.system(), platform.release(), platform.machine()),
                "system": platform.system(),
                "release": platform.release(),
                "machine": platform.machine(),
                "python": sys.version.split()[0],
            },
            "tools": {
                "scrcpy_path": scrcpy or "",
                "scrcpy_version": self._probe.get("scrcpy_version", ""),
                "adb_path": adb or "",
                "adb_version": self._probe.get("adb_version", ""),
            },
            "devices": devices,
            "settings": self.current_settings(),
            "analytics": dict(self.stats),
            "crash_reports": self._list_crash_reports(),
            "events": list(self.logger.events),
            "log": list(self.diag),
        }

    def _list_crash_reports(self):
        try:
            return sorted(n for n in os.listdir(self.crash_reports_dir())
                          if n.startswith("crash-"))
        except Exception:  # noqa: BLE001
            return []

    def save_report(self, rep, path, fmt):
        fmt = (fmt or "txt").lower().lstrip(".")
        t0 = time.time()
        try:
            if fmt == "txt":
                with open(path, "w", encoding="utf-8") as f:
                    f.write(render_txt(rep))
            elif fmt == "md":
                with open(path, "w", encoding="utf-8") as f:
                    f.write(render_md(rep))
            elif fmt == "json":
                with open(path, "w", encoding="utf-8") as f:
                    json.dump(rep, f, indent=2, ensure_ascii=False)
            elif fmt == "jsonl":
                write_jsonl(rep, path)
            elif fmt == "csv":
                write_csv(rep, path)
            elif fmt == "zip":
                write_zip(rep, path, self.logs_dir(), self.crash_reports_dir())
            elif fmt == "xlsx":
                try:
                    write_xlsx(rep, path)
                except ImportError:
                    raise RuntimeError("xlsx ke liye openpyxl chahiye.\n\nInstall: pip install openpyxl\n"
                                       "(ya CSV/TXT/JSONL/ZIP format use karo)")
            else:
                raise RuntimeError("Unknown format: %s" % fmt)
        except RuntimeError:
            self._ev("export_saved", "error", "Diagnostics",
                     props={"format": fmt, "outcome": "failed", "error_code": "export_failed"})
            raise
        except Exception:  # noqa: BLE001
            self._ev("export_saved", "error", "Diagnostics",
                     props={"format": fmt, "outcome": "failed", "error_code": "export_error"})
            raise
        dur = int((time.time() - t0) * 1000)
        self.stats["exports"] += 1
        self._ev("export_saved", "info", "Diagnostics",
                 props={"format": fmt, "outcome": "success", "duration_ms": dur})
        self._ev("performance_marker", "info", "Diagnostics",
                 props={"flow": "export", "duration_ms": dur})

    def email_report(self, gmail=True):
        rep = self.build_report()
        d = self.diagnostics_dir()
        ts = time.strftime("%Y%m%d_%H%M%S")
        path = os.path.join(d, "PhoneMirror_diagnostics_%s.zip" % ts)
        try:
            write_zip(rep, path, self.logs_dir(), self.crash_reports_dir())
        except Exception as exc:  # noqa: BLE001
            self._ev("email_compose", "error", "Diagnostics",
                     props={"outcome": "failed", "error_code": "report_save_failed"})
            raise RuntimeError("Report save nahi hui: %s" % exc)

        body = render_email_body(rep, path)
        subj = "PhoneMirror diagnostics report - %s" % rep["generated_at"]

        try:
            self.root.clipboard_clear()
            self.root.clipboard_append(render_txt(rep))
        except Exception:  # noqa: BLE001
            pass

        if gmail:
            url = ("https://mail.google.com/mail/?view=cm&fs=1&to=%s&su=%s&body=%s"
                   % (quote(DIAG_EMAIL), quote(subj), quote(body)))
        else:
            url = "mailto:%s?subject=%s&body=%s" % (DIAG_EMAIL, quote(subj), quote(body))
        webbrowser.open(url)

        self.stats["emails"] += 1
        self._ev("email_compose", "info", "Diagnostics",
                 props={"outcome": "success", "channel": "gmail" if gmail else "mailto"})
        self._log(None, "[i] Email compose khula -> %s (report: %s)" % (DIAG_EMAIL, path))
        self.status.configure(text="Email compose khula. Report file: %s" % path)
        return path

    def share_bundle(self):
        """Zip bundle banao, uska path clipboard me copy karo, aur folder kholo.
        Phir user kisi bhi app (WhatsApp / Telegram / email) me paste ya attach kar sakta hai."""
        rep = self.build_report()
        d = self.diagnostics_dir()
        ts = time.strftime("%Y%m%d_%H%M%S")
        path = os.path.join(d, "PhoneMirror_diagnostics_%s.zip" % ts)
        write_zip(rep, path, self.logs_dir(), self.crash_reports_dir())
        try:
            self.root.clipboard_clear()
            self.root.clipboard_append(path)
        except Exception:  # noqa: BLE001
            pass
        try:
            if IS_WINDOWS:
                os.startfile(d)  # noqa: S606
        except Exception:  # noqa: BLE001
            pass
        self.stats["shares"] += 1
        self._ev("share_bundle", "info", "Diagnostics",
                 props={"outcome": "success", "format": "zip", "channel": "clipboard+folder"})
        self._log(None, "[i] Share bundle ready: %s (path clipboard me)" % path)
        self.status.configure(text="Share ready - zip path clipboard me, folder khul gayi.")
        return path

    def open_diagnostics(self):
        self._last_screen = "Diagnostics"
        self._last_action = "open_diagnostics"
        self._ev("screen_view", "info", "Diagnostics",
                 props={"screen_name": "Diagnostics", "previous_screen": "Main"})
        self.analytics_snapshot("diagnostics_open")
        self.refresh_probe()
        if self._diag_win is not None and self._diag_win.winfo_exists():
            self._diag_win.lift()
            self._diag_win.refresh()
            return
        self._diag_win = DiagnosticsWindow(self)

    def _tk_exception(self, exc, val, tb):
        """Tkinter callback me uncaught exception -> crash report file + crash_report event.
        Raw traceback sirf local file me jaata hai; event me hash + context."""
        tb_text = "".join(traceback.format_exception(exc, val, tb))
        etype = getattr(exc, "__name__", str(exc))
        path = write_crash_report(etype, val, tb_text,
                                  {"session_id": self.session_id,
                                   "last_screen": self._last_screen,
                                   "last_action": self._last_action},
                                  base=self.diagnostics_dir())
        self.stats["crashes"] += 1
        self._ev("crash_report", "error", "Main", props={
            "fatal": False,
            "exception_type": etype,
            "stack_trace_hash": hashlib.sha256(tb_text.encode("utf-8")).hexdigest()[:16],
            "last_screen": self._last_screen,
            "last_action": self._last_action,
            "report_file": os.path.basename(path) if path else "",
        })
        self._log(None, "[!] Unhandled error: %s: %s (report: %s)"
                  % (etype, val, os.path.basename(path) if path else "-"))

    # ---------------- presets / settings ----------------
    def _apply_saved(self, s):
        """config se saved settings wapas lagao (restart par yaad rahein)."""
        if not isinstance(s, dict):
            return

        def sset(var, key, cast=None):
            if key in s and s[key] is not None:
                try:
                    var.set(cast(s[key]) if cast else s[key])
                except Exception:  # noqa: BLE001
                    pass

        sset(self.video_source, "video_source")
        sset(self.max_size, "max_size", int)
        sset(self.bitrate, "bitrate")
        sset(self.max_fps, "max_fps", int)
        sset(self.codec, "codec")
        sset(self.camera_facing, "camera_facing")
        sset(self.camera_size, "camera_size")
        sset(self.camera_fps, "camera_fps", int)
        sset(self.audio, "audio", bool)
        sset(self.audio_codec, "audio_codec")
        sset(self.audio_bitrate, "audio_bitrate")
        sset(self.audio_buffer, "audio_buffer", int)
        sset(self.record, "record", bool)
        sset(self.record_path, "record_path")
        sset(self.no_control, "no_control", bool)
        sset(self.keyboard_mode, "keyboard_mode")
        sset(self.turn_screen_off, "turn_screen_off", bool)
        sset(self.stay_awake, "stay_awake", bool)
        sset(self.always_on_top, "always_on_top", bool)
        sset(self.print_fps, "print_fps", bool)
        sset(self.auto_reconnect, "auto_reconnect", bool)
        sset(self.auto_refresh, "auto_refresh", bool)
        sset(self.reconnect_delay, "reconnect_delay", int)
        sset(self.reconnect_max, "reconnect_max", int)
        sset(self.connection_mode, "connection_mode")
        sset(self.wifi_ip, "wifi_ip")
        sset(self.wifi_port, "wifi_port", int)
        sset(self.scrcpy_path, "scrcpy_path")
        sset(self.adb_path, "adb_path")
        self._want_serials = list(s.get("selected_serials") or [])
        # derived UI sync
        try:
            self._on_size(None)
            self._on_source()
        except Exception:  # noqa: BLE001
            pass

    def save_settings(self):
        """Saari settings + last device config me save karo."""
        try:
            cfg = load_config()
            cfg["settings"] = {
                "video_source": self.video_source.get(),
                "max_size": self._int_var(self.max_size, 1024),
                "bitrate": self.bitrate.get(),
                "max_fps": self._int_var(self.max_fps, 60),
                "codec": self.codec.get(),
                "camera_facing": self.camera_facing.get(),
                "camera_size": self.camera_size.get(),
                "camera_fps": self._int_var(self.camera_fps, 30),
                "audio": bool(self.audio.get()),
                "audio_codec": self.audio_codec.get(),
                "audio_bitrate": self.audio_bitrate.get(),
                "audio_buffer": self._int_var(self.audio_buffer, 50),
                "record": bool(self.record.get()),
                "record_path": self.record_path.get(),
                "no_control": bool(self.no_control.get()),
                "keyboard_mode": self.keyboard_mode.get(),
                "turn_screen_off": bool(self.turn_screen_off.get()),
                "stay_awake": bool(self.stay_awake.get()),
                "always_on_top": bool(self.always_on_top.get()),
                "print_fps": bool(self.print_fps.get()),
                "auto_reconnect": bool(self.auto_reconnect.get()),
                "auto_refresh": bool(self.auto_refresh.get()),
                "reconnect_delay": self._int_var(self.reconnect_delay, 3),
                "reconnect_max": self._int_var(self.reconnect_max, 0),
                "connection_mode": self.connection_mode.get(),
                "wifi_ip": self.wifi_ip.get(),
                "wifi_port": self._int_var(self.wifi_port, 5555),
                "scrcpy_path": self.scrcpy_path.get(),
                "adb_path": self.adb_path.get(),
                "selected_serials": self.selected_serials(),
            }
            cfg["diagnostics_enabled"] = bool(self.diag_enabled.get())
            save_config(cfg)
        except Exception:  # noqa: BLE001
            pass

    def apply_preset(self, name):
        p = PRESETS[name]
        self.max_size.set(p["max_size"])
        self._on_size(None)
        self.bitrate.set(p["bitrate"])
        self.max_fps.set(p["max_fps"])
        self.codec.set(p["codec"])
        self.audio.set(p["audio"])
        self.audio_codec.set(p["audio_codec"])
        self.audio_bitrate.set(p["audio_bitrate"])
        self.audio_buffer.set(p["audio_buffer"])
        self.preset_desc.configure(text="[%s]  %s" % (name, p["desc"]))
        self.status.configure(text="Preset lagaya: %s" % name)
        self._ev("settings_changed", "info", "Main",
                 props={"setting": "preset", "value": name})
        self.save_settings()

    def _int_var(self, var, default=0, lo=None, hi=None):
        """IntVar/Entry se safe integer nikalo.
        User ne non-numeric text type kiya ho to crash nahi, default use hota hai."""
        try:
            v = int(var.get())
        except Exception:  # noqa: BLE001
            v = default
        if lo is not None and v < lo:
            v = lo
        if hi is not None and v > hi:
            v = hi
        return v

    def current_settings(self):
        return {
            "video_source": self.video_source.get(),
            "max_size": self._int_var(self.max_size, 1024, 240, 4096),
            "bitrate": self.bitrate.get(),
            "max_fps": self._int_var(self.max_fps, 60, 1, 240),
            "codec": self.codec.get(),
            "camera_facing": self.camera_facing.get(),
            "camera_size": self.camera_size.get(),
            "camera_fps": self._int_var(self.camera_fps, 30, 1, 120),
            "audio": self.audio.get(),
            "audio_codec": self.audio_codec.get(),
            "audio_bitrate": self.audio_bitrate.get(),
            "audio_buffer": self._int_var(self.audio_buffer, 50, 0, 1000),
            "record": self.record.get(),
            "record_path": redact(self.record_path.get()),
            "no_control": self.no_control.get(),
            "keyboard_mode": self.keyboard_mode.get(),
            "turn_screen_off": self.turn_screen_off.get(),
            "stay_awake": self.stay_awake.get(),
            "always_on_top": self.always_on_top.get(),
            "print_fps": self.print_fps.get(),
            "auto_reconnect": self.auto_reconnect.get(),
            "auto_refresh": self.auto_refresh.get(),
            "reconnect_delay": self._int_var(self.reconnect_delay, 3, 1, 600),
            "reconnect_max": self._int_var(self.reconnect_max, 0, 0, 100000),
            "connection_mode": self.connection_mode.get(),
        }

    def selected_serials(self):
        out = []
        for i in self.device_list.curselection():
            if i < len(self.devices):
                out.append(self.devices[i]["serial"])
        return out

    def primary_serial(self):
        s = self.selected_serials()
        return s[0] if s else None

    def record_path_for(self, serial):
        safe = re.sub(r"[^A-Za-z0-9_.-]", "_", serial or "device")
        ts = time.strftime("%Y%m%d_%H%M%S")
        base = self.record_path.get().strip()
        if base:
            root_, ext = os.path.splitext(base)
            ext = ext or ".mp4"
            if self._launch_multi:
                return "%s_%s%s" % (root_, safe, ext)
            return base
        vids = os.path.join(os.path.expanduser("~"), "Videos")
        if not os.path.isdir(vids):
            vids = os.path.expanduser("~")
        return os.path.join(vids, "PhoneMirror_%s_%s.mp4" % (safe, ts))

    # ---------------- Wi-Fi helpers ----------------
    def detect_ip(self):
        adb = self._adb()
        if not adb:
            self._log(None, "[!] adb nahi mila.")
            self._ev("error_state_shown", "error", "Main",
                     props={"error_code": "adb_missing", "message_key": "adb_not_found"})
            return
        self.status.configure(text="Phone IP dhoondh rahe hain...")
        serial = self.primary_serial()

        def done(ok, res):
            ip = res if ok else None
            if ip:
                self.wifi_ip.set(ip)
                self._log(None, "[i] Phone IP: %s" % ip)
                self.status.configure(text="Phone IP: %s" % ip)
            else:
                self._log(None, "[!] IP detect nahi hua.")
                self._ev("error_state_shown", "error", "Main",
                         props={"error_code": "ip_not_found", "message_key": "ip_not_detected"})
                self.toast("IP detect nahi hua - phone Wi-Fi par hai aur USB se juda hai?", "warn")

        self._run_bg(lambda: detect_phone_ip(adb, serial), done)

    def enable_wifi(self):
        adb = self._adb()
        if not adb:
            self._log(None, "[!] adb nahi mila.")
            return
        serial = self.primary_serial()
        if not serial:
            self._ev("error_state_shown", "warn", "Main",
                     props={"error_code": "no_device", "message_key": "no_device_selected"})
            self.toast("Pehle USB se phone jodo aur Refresh dabao.", "warn")
            return
        port = self._int_var(self.wifi_port, 5555, 1, 65535)
        self.status.configure(text="Wi-Fi mode on kar rahe hain...")

        def done(ok, res):
            if not ok:
                self._log(None, "[!] adb tcpip error: %s" % res)
                return
            self._log(None, "[i] adb tcpip %d -> %s" % (port, "OK" if res["ok"] else res["out"]))
            if res["ok"]:
                self.status.configure(text="Wi-Fi mode on. Cable nikaalo aur 'Connect' dabao.")
            else:
                self._ev("error_state_shown", "error", "Main",
                         props={"error_code": "tcpip_failed", "message_key": "enable_wifi_failed"})
                self.toast("Enable Wi-Fi fail: %s" % (res["out"] or "adb tcpip fail hua"), "error")

        self._adb_bg(["-s", serial, "tcpip", str(port)], "Main", done)

    def connect_wifi(self):
        adb = self._adb()
        if not adb:
            self._log(None, "[!] adb nahi mila.")
            return
        ip = self.wifi_ip.get().strip()
        if not ip:
            self._ev("error_state_shown", "warn", "Main",
                     props={"error_code": "no_ip", "message_key": "ip_missing"})
            self.toast("Phone IP daalo ya 'Auto-detect IP' dabao.", "warn")
            return
        target = "%s:%d" % (ip, self._int_var(self.wifi_port, 5555, 1, 65535))
        self.status.configure(text="Connect kar rahe hain...")

        def done(ok, res):
            if not ok:
                self._log(None, "[!] adb connect error: %s" % res)
                return
            self._log(None, "[i] adb connect %s -> %s" % (target, res["out"]))
            self.refresh_devices()
            self.status.configure(text="Connect try kiya. Device list check karo.")

        self._adb_bg(["connect", target], "Main", done)

    def disconnect_wifi(self):
        adb = self._adb()
        ip = self.wifi_ip.get().strip()
        if not (adb and ip):
            return
        target = "%s:%d" % (ip, self._int_var(self.wifi_port, 5555, 1, 65535))

        def done(ok, res):
            if not ok:
                self._log(None, "[!] adb disconnect error: %s" % res)
                return
            self._log(None, "[i] adb disconnect %s -> %s" % (target, res["out"]))
            self.refresh_devices()

        self._adb_bg(["disconnect", target], "Main", done)

    # ---------------- devices ----------------
    def refresh_devices(self, quiet=False):
        """Devices ko background me scan karo - UI block nahi hoti.
        quiet=True (auto-refresh): UI/log sirf tab update hote hain jab device set badla ho."""
        adb = self._adb()
        if not adb:
            self.device_status.configure(text="adb nahi mila. Upar path set karo.", foreground="#b00")
            self._log(None, "[!] adb nahi mila.")
            self._ev("device_state_resolved", "error", "Main",
                     props={"state": "unknown", "device_count": 0, "error_code": "adb_missing"})
            return
        if self._refreshing:
            return
        self._refreshing = True
        t0 = time.time()
        if not quiet:
            self.device_status.configure(text="Devices dhoondh rahe hain...", foreground="#666")

        prev = set()
        for i in self.device_list.curselection():
            if i < len(self.devices):
                prev.add(self.devices[i]["serial"])
        if not prev:
            prev = set(self._want_serials or [])   # saved selection restore
        sig_old = self._dev_sig()

        def done(ok, res):
            self._refreshing = False
            if not ok:
                if not quiet:
                    self.device_status.configure(text="Device list error: %s" % res, foreground="#b00")
                    self._log(None, "[!] %s" % res)
                    self._ev("device_refresh_result", "error", "Main",
                             props={"outcome": "failed", "error_code": "list_failed"})
                return
            newdevs, err = res
            if err:
                if not quiet:
                    self.device_status.configure(text="Device list error: %s" % err, foreground="#b00")
                    self._log(None, "[!] %s" % err)
                    self._ev("device_refresh_result", "error", "Main",
                             props={"outcome": "failed", "error_code": "list_failed"})
                return

            if quiet and self._dev_sig(newdevs) == sig_old:
                return                      # kuch nahi badla -> UI/log churn nahi

            self.devices = newdevs
            self.device_list.delete(0, "end")
            for d in self.devices:
                self.device_list.insert("end", d["label"])

            restored = False
            for i, d in enumerate(self.devices):
                if d["serial"] in prev:
                    self.device_list.selection_set(i)
                    restored = True
            if not restored and self.devices:
                self.device_list.selection_set(0)

            ready = [d for d in self.devices if d["state"] == "device"]
            if ready:
                self.device_status.configure(text="%d device ready. Multi-select kar sakte ho." % len(ready),
                                             foreground="#080")
            elif self.devices:
                self.device_status.configure(text="Device unauthorized/offline. USB debugging allow karo.",
                                             foreground="#b60")
            else:
                self.device_status.configure(text="Koi device nahi mila.", foreground="#b60")
            self._log(None, "[i] Devices refresh: %d mile." % len(self.devices))

            if ready:
                state = "authorized"
            elif self.devices:
                state = "unauthorized"
            else:
                state = "none"
            self._ev("device_state_resolved", "info", "Main",
                     props={"state": state, "device_count": len(self.devices), "ready_count": len(ready)})
            self._ev("device_refresh_result", "info", "Main",
                     props={"outcome": "success", "device_count": len(self.devices), "ready_count": len(ready)})
            self.stats["devices_seen"] = max(self.stats["devices_seen"], len(self.devices))
            self._ev("performance_marker", "info", "Main",
                     props={"flow": "device_scan", "duration_ms": int((time.time() - t0) * 1000)})
            if state == "unauthorized":
                # security: phone ne USB debugging permission nahi di
                self._ev("device_authorization_failed", "warn", "Main",
                         props={"state": state, "device_count": len(self.devices)})

        self._run_bg(lambda: list_devices(adb), done)

    def _dev_sig(self, devs=None):
        devs = self.devices if devs is None else devs
        return tuple(sorted((d.get("serial", ""), d.get("state", "")) for d in devs))

    def _auto_refresh_tick(self):
        """Har 5s me quiet refresh - naya phone khud dikh jaye."""
        try:
            if self.auto_refresh.get() and not self._refreshing:
                self.refresh_devices(quiet=True)
        except Exception:  # noqa: BLE001
            pass
        try:
            self.root.after(5000, self._auto_refresh_tick)
        except Exception:  # noqa: BLE001
            pass

    # ---------------- run ----------------
    def copy_command(self):
        scrcpy = self.scrcpy_path.get().strip() or "scrcpy"
        serials = self.selected_serials()
        if not serials:
            self.status.configure(text="Pehle device select karo.")
            return
        lines = []
        for serial in serials:
            s = self.current_settings()
            if s.get("record"):
                s["record_path"] = self.record_path_for(serial)
            cmd = build_command(scrcpy, serial, s)
            lines.append(" ".join('"%s"' % c if " " in c else c for c in cmd))
        self.root.clipboard_clear()
        self.root.clipboard_append("\n".join(lines))
        self.status.configure(text="%d command clipboard me copy hue." % len(lines))

    def _resolve_scrcpy(self):
        scrcpy = self.scrcpy_path.get().strip()
        if not scrcpy or not os.path.exists(scrcpy):
            scrcpy = find_executable("scrcpy")
            if scrcpy:
                self.scrcpy_path.set(scrcpy)
        return scrcpy

    def launch(self):
        scrcpy = self._resolve_scrcpy()
        if not scrcpy:
            self._ev("error_state_shown", "error", "Main",
                     props={"error_code": "scrcpy_missing", "message_key": "scrcpy_not_found"})
            self.toast("scrcpy nahi mila - path set karo (winget install Genymobile.scrcpy)",
                       "error", ms=9000)
            return
        serials = self.selected_serials()
        if not serials:
            self._ev("error_state_shown", "warn", "Main",
                     props={"error_code": "no_device", "message_key": "no_device_selected"})
            self.toast("Pehle device select karo (Ctrl/Shift + click).", "warn")
            return

        self._launch_attempt = self.logger.new_attempt_id()
        self._launch_multi = len(serials) > 1
        self._last_action = "launch_mirroring"
        self.stats["mirrors_started"] += len(serials)
        self.save_settings()
        self._ev("mirror_requested", "info", "Main", attempt_id=self._launch_attempt,
                 props={"device_count": len(serials), "video_source": self.video_source.get(),
                        "connection_mode": self.connection_mode.get()})
        for serial in serials:
            sess = self.sessions.get(serial) or Session(serial)
            self.sessions[serial] = sess
            sess.user_stopped = False
            sess.reconnect_attempts = 0
            self._start_session(serial)
        self.status.configure(text="%d session start ho raha hai..." % len(serials))
        self._refresh_buttons()

    def _start_session(self, serial):
        sess = self.sessions.get(serial) or Session(serial)
        self.sessions[serial] = sess

        scrcpy = self._resolve_scrcpy()
        if not scrcpy:
            self._log(serial, "[!] scrcpy nahi mila, rok raha hoon.")
            self.stats["mirrors_failed"] += 1
            self._ev("mirror_result", "error", "Main", attempt_id=self._launch_attempt,
                     props={"outcome": "failed", "error_code": "scrcpy_missing",
                            "device": self._dev_pseudonym(serial)})
            sess.running_intent = False
            self._refresh_buttons()
            return

        s = self.current_settings()
        if s.get("record"):
            s["record_path"] = self.record_path_for(serial)
        cmd = build_command(scrcpy, serial, s)
        self._log(serial, "$ " + " ".join(cmd))

        try:
            sess.proc = subprocess.Popen(
                cmd, stdout=subprocess.PIPE, stderr=subprocess.STDOUT,
                text=True, bufsize=1, creationflags=CREATE_NO_WINDOW,
            )
        except Exception as exc:  # noqa: BLE001
            self._log(serial, "[!] Launch fail: %s" % exc)
            self.stats["mirrors_failed"] += 1
            self._ev("mirror_result", "error", "Main", attempt_id=self._launch_attempt,
                     props={"outcome": "failed", "error_code": "launch_failed",
                            "device": self._dev_pseudonym(serial)})
            sess.proc = None
            sess.running_intent = False
            self._refresh_buttons()
            return

        sess.start_time = time.time()
        sess.running_intent = True
        self.stats["mirrors_succeeded"] += 1
        self._ev("mirror_result", "info", "Main", attempt_id=self._launch_attempt,
                 props={"outcome": "success", "device": self._dev_pseudonym(serial)})
        threading.Thread(target=self._reader, args=(serial, sess.proc), daemon=True).start()
        self._refresh_buttons()

    def _reader(self, serial, proc):
        try:
            for line in iter(proc.stdout.readline, ""):
                if not line:
                    break
                self.log_q.put(("line", serial, line.rstrip()))
        finally:
            self.log_q.put(("end", serial))

    def _poll_log(self):
        """Main-thread pump: log lines, session ends, aur background callbacks.
        Har cycle me limited items process hote hain taaki UI responsive rahe."""
        processed = 0
        try:
            while processed < 400:
                item = self.log_q.get_nowait()
                processed += 1
                if item[0] == "line":
                    self._log(item[1], item[2])
                elif item[0] == "end":
                    self._on_session_end(item[1])
                elif item[0] == "call":
                    try:
                        item[1]()
                    except Exception:  # noqa: BLE001
                        pass
        except queue.Empty:
            pass
        if processed >= 200:
            # debug_trace: bada backlog ek cycle me nikla (UI responsiveness check)
            self._ev("queue_drain", "trace", "Main", props={"processed": processed})
        self.root.after(50, self._poll_log)

    def _on_session_end(self, serial):
        sess = self.sessions.get(serial)
        if sess is None:
            return
        if sess.proc is not None:
            elapsed = time.time() - sess.start_time
            self._log(serial, "[i] band ho gaya (%.0f s)." % elapsed)
            sess.proc = None
            if elapsed >= 15:
                sess.reconnect_attempts = 0

        if sess.running_intent and not sess.user_stopped and self.auto_reconnect.get():
            delay = self._int_var(self.reconnect_delay, 3, 1, 600)
            self._log(serial, "[i] Auto-reconnect %d s me..." % delay)
            self.root.after(delay * 1000, lambda: self._attempt_reconnect(serial))
        else:
            sess.running_intent = False
        self._refresh_buttons()

    def _attempt_reconnect(self, serial):
        sess = self.sessions.get(serial)
        if sess is None or not sess.running_intent or sess.user_stopped or not self.auto_reconnect.get():
            return
        maxa = self._int_var(self.reconnect_max, 0, 0, 100000)
        if maxa and sess.reconnect_attempts >= maxa:
            self._log(serial, "[!] Max attempts (%d) pura." % maxa)
            sess.running_intent = False
            self._refresh_buttons()
            return

        adb = self._adb()
        if not adb:
            self._log(serial, "[!] adb nahi mila, reconnect rok raha hoon.")
            sess.running_intent = False
            self._refresh_buttons()
            return

        def work():
            devs, _ = list_devices(adb)
            return any(d["serial"] == serial and d["state"] == "device" for d in devs)

        def done(ok, present):
            sess2 = self.sessions.get(serial)
            if (sess2 is None or not sess2.running_intent
                    or sess2.user_stopped or not self.auto_reconnect.get()):
                return
            if ok and present:
                sess2.reconnect_attempts += 1
                self._log(serial, "[i] Reconnect attempt #%d" % sess2.reconnect_attempts)
                att = self.logger.new_attempt_id("rc")
                self._ev("reconnect_attempt", "info", "Main", attempt_id=att,
                         props={"attempt_number": sess2.reconnect_attempts,
                                "device": self._dev_pseudonym(serial)})
                # user-friendliness signal: reconnect_loop
                if sess2.reconnect_attempts >= 3:
                    self._ev("reconnect_loop", "warn", "Main", attempt_id=att,
                             props={"attempt_count": sess2.reconnect_attempts,
                                    "device": self._dev_pseudonym(serial)})
                self.stats["reconnects"] += 1
                self._start_session(serial)
                self._ev("reconnect_result", "info", "Main", attempt_id=att,
                         props={"outcome": "success"})
            else:
                delay = self._int_var(self.reconnect_delay, 3, 1, 600)
                self._log(serial, "[i] Device nahi mila, %d s baad try..." % delay)
                self.root.after(delay * 1000, lambda: self._attempt_reconnect(serial))

        self._run_bg(work, done)

    def _refresh_buttons(self):
        active = any(s.active() for s in self.sessions.values())
        self.launch_btn.configure(state="disabled" if active else "normal")
        self.stop_btn.configure(state="normal" if active else "disabled")

    def stop(self):
        for sess in self.sessions.values():
            sess.user_stopped = True
            sess.running_intent = False
            if sess.proc is not None:
                try:
                    sess.proc.terminate()
                except Exception:  # noqa: BLE001
                    pass
        self._log(None, "[i] Stop dabaya, sab sessions band ho rahe hain.")
        self.status.configure(text="Ruk raha hai...")

    def on_close(self):
        for sess in self.sessions.values():
            sess.user_stopped = True
            sess.running_intent = False
            if sess.proc is not None:
                try:
                    sess.proc.terminate()
                except Exception:  # noqa: BLE001
                    pass
        self._ev("session_end", "info", "Main", props={"reason": "user_close"})
        self.analytics_snapshot("session_end")
        self.save_settings()
        try:
            self.filelog.close()
        except Exception:  # noqa: BLE001
            pass
        self.root.destroy()


def main():
    root = tk.Tk()
    try:
        ttk.Style().theme_use("vista" if IS_WINDOWS else "clam")
    except tk.TclError:
        pass
    try:
        app = PhoneMirrorApp(root)
    except Exception:  # noqa: BLE001
        # Startup crash: crash report file + user ko dikhao (silent death se bacho)
        tb = traceback.format_exc()
        path = write_crash_report("StartupError", "app init fail", tb,
                                  {"phase": "PhoneMirrorApp.__init__"})
        try:
            messagebox.showerror("PhoneMirror start nahi ho paya",
                                 "Startup error:\n\n%s%s"
                                 % (tb[-900:], ("\n\nReport: %s" % path) if path else ""))
        except Exception:  # noqa: BLE001
            pass
        try:
            root.destroy()
        except Exception:  # noqa: BLE001
            pass
        return
    root.report_callback_exception = app._tk_exception
    root.protocol("WM_DELETE_WINDOW", app.on_close)
    root.mainloop()


if __name__ == "__main__":
    main()
